# Folha única MedSum · estilo 2 (regra da Laura, 22/10/2026)
# Três layouts leves, sempre com "Não confunda" e "Mnemônicos":
#   A = arejada (texto em tópicos + coluna lateral com quadros)
#   B = com fluxograma (fluxo da conduta em destaque + cartões)
#   C = visual (faixa bege no topo, desenho/cartões de números, passo a passo horizontal)
# Uso: render(nome, dict(layout='A'|'B'|'C', titulo, sub, corpo=HTML, fontes)). O corpo usa as classes do CSS abaixo.
# Saída: PDF (A4, 1 página), PNG e Word (página inteira como imagem) em saida/<PASTA>/.
import os, re, subprocess, html as H

AQUI = os.path.dirname(os.path.abspath(__file__))
SAIDA = os.path.join(AQUI, 'saida', os.environ.get('PASTA', 'folhas')); os.makedirs(SAIDA, exist_ok=True)
LOGO = os.environ.get('LOGO', '/home/user/medsum/marca/logo-ms-vinho.png')

CSS = """@page{size:A4;margin:0}*{box-sizing:border-box;margin:0;padding:0}
:root{--ou:#8a5f12;--ou2:#a8771c;--ous:#f8f0de;--vi:#921827;--vis:#f7e8ea;--ink:#2b2224;--mu:#76696b;--ln:#ece3d6;--ok:#eef6f0;--okl:#cfe5d5}
html,body{width:210mm;height:297mm}
body{font-family:Inter,sans-serif;font-weight:400;color:var(--ink);font-size:var(--fs,8.2pt);line-height:var(--lh,1.52);position:relative;overflow:hidden;-webkit-font-smoothing:antialiased}
b,strong{font-weight:600}.k{color:var(--vi);font-weight:600}.m{color:var(--mu)}
.selo{display:inline-block;border:1px solid var(--ou2);color:var(--ou);font-weight:600;letter-spacing:.16em;font-size:6.2pt;padding:2px 9px;border-radius:99px}
h1{font-family:'DejaVu Serif',serif;font-weight:400;color:var(--ou);letter-spacing:-.01em;font-size:var(--h1,24pt);line-height:1.1;margin:4px 0 1px}
.sub{color:var(--mu);font-size:8.2pt}
header{display:flex;justify-content:space-between;align-items:flex-end}
header img{height:12mm}
.L-A header,.L-B header{padding:10mm 12mm 3mm;margin:0 0 4mm;border-bottom:1pt solid var(--ou2);margin-left:12mm;margin-right:12mm;padding-left:0;padding-right:0}
.L-C header{background:var(--ous);padding:8mm 12mm 5mm;margin-bottom:5mm}
main{padding:0 12mm}
h2{font-family:'DejaVu Serif',serif;font-weight:400;font-size:11.5pt;color:var(--ou);margin:0 0 1.6mm;padding-bottom:1px;border-bottom:.5pt solid var(--ln)}
h2.t{font-family:Inter,sans-serif;font-size:6.9pt;letter-spacing:.14em;text-transform:uppercase;font-weight:600;border:0;margin-bottom:1.6mm}
.sec{margin-bottom:var(--gap,3.6mm)}
p{margin:0 0 2px}
ul{padding-left:11px;margin:0}li{margin-bottom:1.6px}li::marker{color:var(--ou2)}
.g2{display:grid;grid-template-columns:1fr 1fr;gap:6mm}
.g3{display:grid;grid-template-columns:1fr 1fr 1fr;gap:4mm}
.gL{display:grid;grid-template-columns:1fr 60mm;gap:7mm}
.gB{display:grid;grid-template-columns:1fr 1fr;gap:6mm}
.c{border:1px solid var(--ln);border-radius:10px;padding:3mm 3.8mm;margin-bottom:var(--gap,3.6mm);background:#fff}
.box{background:var(--ous);border-radius:10px;padding:3.6mm 4mm;margin-bottom:var(--gap,3.6mm)}
.box h3,.nc h3,.mn h3{font-family:'DejaVu Serif',serif;font-weight:400;font-size:10.5pt;color:var(--ou);margin-bottom:1.8mm}
.nc{background:var(--vis);border-radius:10px;padding:3.6mm 4mm;margin-bottom:var(--gap,3.6mm)}
.nc h3{color:var(--vi)}
.nc li{margin-bottom:2.4px}
.mn{border:1.2pt dashed var(--ou2);border-radius:10px;padding:3.2mm 4mm;margin-bottom:var(--gap,3.6mm);background:#fffdf8}
.mn h3::before{content:'💡 ';font-size:9pt}
.mn .it{margin-bottom:2.2mm}.mn .it:last-child{margin-bottom:0}
.mn .it>b{display:block;color:var(--vi);font-family:'DejaVu Serif',serif;font-weight:400;font-size:10pt;letter-spacing:.02em}
.mn .it>span{color:var(--mu);font-size:.95em}
table{width:100%;border-collapse:collapse;margin:1px 0 2px;font-size:.96em}
th{font-weight:600;color:var(--ou);text-align:left;font-size:.86em;letter-spacing:.06em;text-transform:uppercase;padding:2px 4px;border-bottom:1pt solid var(--ou2)}
td{padding:2px 4px;border-bottom:.5pt solid var(--ln);vertical-align:top}
td:first-child{font-weight:600}
/* fluxograma (B) */
.flx{display:flex;flex-direction:column;text-align:center;font-size:.93em;line-height:1.35}
.fb{border:1px solid var(--ln);background:#fff;border-radius:8px;padding:4px 6px}
.fb.esc{background:var(--ink);color:#fff;border:0;font-weight:600}
.fb.vi{background:var(--vi);color:#fff;border:0;font-weight:600}
.fb.ok{background:var(--ok);border-color:var(--okl)}
.fb.ou{background:var(--ous);border-color:#ead9b4}
.fa{color:var(--ou2);font-weight:600;font-size:.9em;line-height:1.15;margin:2px 0}
.fr{display:grid;grid-template-columns:1fr 1fr;gap:5px}.fr3{display:grid;grid-template-columns:1fr 1fr 1fr;gap:5px}
.fp{font-size:.92em;color:var(--mu);padding:2px 1px 0}
/* cartões de número e passo a passo (C) */
.cards{display:grid;grid-template-columns:repeat(var(--nc,3),1fr);gap:3.5mm}
.cd{border-radius:10px;padding:3mm 3.5mm;background:#fff;border:1px solid var(--ln)}
.cd .t{font-size:6.5pt;letter-spacing:.13em;text-transform:uppercase;color:var(--ou);font-weight:600;display:block;margin-bottom:1mm}
.cd .n{font-family:'DejaVu Serif',serif;font-size:14pt;color:var(--vi);line-height:1.1;display:block}
.hz{display:grid;grid-template-columns:repeat(var(--np,4),1fr);gap:3.5mm}
.st{position:relative;border-radius:10px;padding:3mm 3.2mm;background:#fff;border:1px solid var(--ln)}
.st i{font-style:normal;display:inline-grid;place-items:center;width:14px;height:14px;border-radius:50%;background:var(--vi);color:#fff;font-size:6.8pt;font-weight:600;margin-bottom:1mm}
.st b{display:block}
.st:not(:last-child)::after{content:'›';position:absolute;right:-2.7mm;top:40%;color:var(--ou2);font-size:12pt}
.vis{display:grid;grid-template-columns:var(--vw,34mm) 1fr;gap:5mm;align-items:center}
footer{position:absolute;left:12mm;right:12mm;bottom:6mm;display:flex;justify-content:space-between;gap:8mm;font-size:6.2pt;color:var(--mu);border-top:.5pt solid var(--ln);padding-top:3px}
footer b{color:var(--vi);letter-spacing:.1em;font-weight:600;white-space:nowrap}
"""


def render(nome, f):
    head = f'''<header><div><span class="selo">REVISE PARA PROVA · RESUMO EM UMA FOLHA</span><h1>{H.escape(f['titulo'])}</h1><p class="sub">{f['sub']}</p></div><img src="file://{LOGO}" alt=""></header>'''
    doc = f'''<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><style>{CSS}body{{{f.get('vars','')}}}{f.get('css','')}</style></head>
<body class="L-{f['layout']}">{head}<main>{f['corpo']}</main>
<footer><span>{f['fontes']}</span><b>MEDSUM · @medsum.oficial</b></footer></body></html>'''
    src = os.path.join(SAIDA, nome + '.html'); open(src, 'w').write(doc)
    pdf, png = os.path.join(SAIDA, nome + '.pdf'), os.path.join(SAIDA, nome + '.png')
    js = f"""const {{chromium}}=require('playwright');(async()=>{{const b=await chromium.launch({{executablePath:'/opt/pw-browsers/chromium'}});
const p=await b.newPage({{viewport:{{width:794,height:1123}},deviceScaleFactor:3}});await p.goto('file://{src}');await p.waitForTimeout(300);
const r=await p.evaluate(()=>{{const m=document.querySelector('main').getBoundingClientRect(),f=document.querySelector('footer').getBoundingClientRect();
 let fundo=0;document.querySelectorAll('main *').forEach(e=>{{const q=e.getBoundingClientRect();if(q.height)fundo=Math.max(fundo,q.bottom);}});return Math.round(fundo-f.top);}});
console.log('sobra(px, negativo=cabe):',r);
await p.pdf({{path:'{pdf}',format:'A4',printBackground:true,margin:{{top:0,right:0,bottom:0,left:0}}}});
await p.screenshot({{path:'{png}'}});await b.close();}})();"""
    jsf = os.path.join(SAIDA, '_r.js'); open(jsf, 'w').write(js)
    env = dict(os.environ, NODE_PATH=subprocess.run(['npm', 'root', '-g'], capture_output=True, text=True).stdout.strip())
    out = subprocess.run(['node', jsf], capture_output=True, text=True, env=env); print(nome, out.stdout.strip(), out.stderr[-300:])
    os.remove(jsf)
    import docx
    from docx.shared import Cm
    d = docx.Document(); s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    for a in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(s, a, Cm(0))
    s.header_distance = s.footer_distance = Cm(0)
    p = d.add_paragraph(); p.paragraph_format.space_after = 0; p.paragraph_format.space_before = 0
    p.add_run().add_picture(png, width=Cm(21.0))
    d.save(os.path.join(SAIDA, nome + '.docx'))
    n = subprocess.run(['pdfinfo', pdf], capture_output=True, text=True).stdout
    print('  páginas:', re.search(r'Pages:\s+(\d+)', n).group(1))


# blocos prontos
def ul(itens): return '<ul>' + ''.join(f'<li>{x}</li>' for x in itens) + '</ul>'
def nc(itens, titulo='Não confunda na prova'): return f'<div class="nc"><h3>{titulo}</h3>{ul(itens)}</div>'
def mn(itens, titulo='Para não esquecer'):
    return f'<div class="mn"><h3>{titulo}</h3>' + ''.join(f'<div class="it"><b>{a}</b><span>{b}</span></div>' for a, b in itens) + '</div>'
def sec(t, corpo, cls='sec'): return f'<div class="{cls}"><h2>{t}</h2>{corpo}</div>'
def card(t, corpo): return f'<div class="c"><h2 class="t">{t}</h2>{corpo}</div>'
def box(t, corpo): return f'<div class="box"><h3>{t}</h3>{corpo}</div>'
def cards(lst): return '<div class="cards" style="--nc:%d">' % len(lst) + ''.join(f'<div class="cd"><span class="t">{t}</span><span class="n">{n}</span>{x}</div>' for t, n, x in lst) + '</div>'
def passos(lst): return '<div class="hz" style="--np:%d">' % len(lst) + ''.join(f'<div class="st"><i>{i+1}</i><b>{t}</b>{x}</div>' for i, (t, x) in enumerate(lst)) + '</div>'
