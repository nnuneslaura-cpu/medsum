# Folha única (A4) no padrão MedSum: HTML -> PDF (Playwright) e Word (página inteira como imagem, igual ao modelo do Diabetes)
import os, subprocess, json, html as H, re

AQUI = os.path.dirname(os.path.abspath(__file__))
SAIDA = os.path.join(AQUI, 'saida', os.environ.get('PASTA','folhas')); os.makedirs(SAIDA, exist_ok=True)
LOGO = '/home/user/medsum/marca/logo-ms-vinho.png'

CSS = """
@page{size:A4;margin:0}
*{box-sizing:border-box;margin:0;padding:0}
:root{--ou:#8a5f12;--ou2:#a8771c;--ous:#f6ecd6;--vi:#921827;--ink:#241a1c;--mu:#6c5f61;--ln:#e6dccb}
html,body{width:210mm;height:297mm}
body{font-family:Inter,'DejaVu Sans',sans-serif;color:var(--ink);font-size:7.55pt;line-height:1.34;padding:7mm 8mm 6mm;position:relative;overflow:hidden}
b{font-weight:700}
.selo{display:inline-block;background:var(--ou2);color:#fff;font-weight:800;letter-spacing:.14em;font-size:6.6pt;padding:2px 8px;border-radius:99px;margin-bottom:3px}
header{display:flex;justify-content:space-between;align-items:flex-start;border-bottom:1.4pt solid var(--ou2);padding-bottom:3px}
h1{font-family:'DejaVu Serif',serif;font-variant:small-caps;font-size:23pt;line-height:1;color:var(--ou);letter-spacing:-.01em}
.sub{color:var(--mu);font-size:8pt;margin-top:2px}
header img{height:15mm}
.mais{display:grid;grid-template-columns:auto repeat(4,1fr);gap:4px;align-items:center;margin:4px 0 4px}
.mais>span{font-weight:800;font-size:6.6pt;letter-spacing:.1em;color:var(--mu);line-height:1.1}
.mais div{border:1px solid var(--ln);border-radius:9px;padding:2px 6px;font-size:6.9pt;line-height:1.25}
.mais div b{color:var(--ou)}
.grid{display:grid;grid-template-columns:1.58fr 1fr;gap:5mm}
.sec{margin-bottom:3.2mm}
.sec h2{display:flex;align-items:center;gap:5px;font-family:'DejaVu Serif',serif;font-variant:small-caps;font-size:12.5pt;color:var(--ou);margin-bottom:2px}
.sec h2 i{font-style:normal;font-family:Inter,sans-serif;font-variant:normal;font-size:8pt;background:var(--ou2);color:#fff;width:15px;height:15px;border-radius:50%;display:grid;place-items:center;flex:none}
.sec h2 small{font-family:Inter,sans-serif;font-variant:normal;font-size:7pt;color:var(--mu);font-weight:400}
h3{font-size:7pt;letter-spacing:.08em;text-transform:uppercase;color:var(--ou);margin:3px 0 1px}
table{width:100%;border-collapse:collapse;margin:1px 0 2px}
th{background:var(--ou2);color:#fff;text-align:left;font-size:6.9pt;padding:2px 4px}
td{padding:1.6px 4px;border-bottom:.6pt solid var(--ln);vertical-align:top}
tr:nth-child(even) td{background:#fbf6ea}
td:first-child{font-weight:700}
.cols{display:grid;gap:4px}
.c2{grid-template-columns:1fr 1fr}.c3{grid-template-columns:1fr 1fr 1fr}
.cx{border:1px solid var(--ln);border-radius:7px;padding:3px 5px}
.cx b.t{display:block;color:var(--ou);font-size:7.1pt;letter-spacing:.04em}
.dest{background:var(--ous);border-left:2.4pt solid var(--ou2);border-radius:4px;padding:3px 6px;margin:2px 0}
ul{padding-left:10px}li{margin:0 0 1px}
.lado{background:#fffaf0;border:1.2pt solid var(--ou2);border-radius:10px;overflow:hidden;align-self:start}
.lado .cab{background:var(--vi);color:#fff;padding:4px 8px}
.lado .cab b{font-family:'DejaVu Serif',serif;font-variant:small-caps;font-size:12pt;display:block}
.lado .cab span{font-size:6.8pt;opacity:.85}
.lado ol{list-style:none;padding:4px 7px 5px;counter-reset:n}
.lado li{counter-increment:n;position:relative;padding:2.5px 0 2.5px 17px;border-bottom:.6pt solid var(--ln)}
.lado li:last-child{border-bottom:0}
.lado li::before{content:counter(n);position:absolute;left:0;top:3px;width:12px;height:12px;border-radius:50%;border:1pt solid var(--vi);color:var(--vi);font-weight:800;font-size:6.2pt;display:grid;place-items:center}
.lado li b{color:var(--vi)}
.conta{margin-top:3mm;border:1.2pt solid var(--ou2);border-radius:10px;padding:4px 8px}
.conta b.t{font-family:'DejaVu Serif',serif;font-variant:small-caps;font-size:11pt;color:var(--ou);display:block;margin-bottom:1px}
footer{position:absolute;left:8mm;right:8mm;bottom:5mm;display:flex;justify-content:space-between;font-size:6.3pt;color:var(--mu);border-top:.6pt solid var(--ln);padding-top:2px}
footer b{color:var(--vi);letter-spacing:.08em}
"""


def md(t):
    t = H.escape(t)
    return re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', t)


def render(nome, f):
    """f: dict com titulo, sub, mais[4], secoes[(titulo, nota, html)], lado_tit, lado_sub, lado[], conta_tit, conta_html, fontes"""
    mais = ''.join(f'<div><b>{i+1}º</b> {md(x)}</div>' for i, x in enumerate(f['mais']))
    secs = ''.join(f'<div class="sec"><h2><i>{i+1}</i>{H.escape(t)}{f" <small>{H.escape(n)}</small>" if n else ""}</h2>{h}</div>'
                   for i, (t, n, h) in enumerate(f['secoes']))
    lado = ''.join(f'<li>{md(x)}</li>' for x in f['lado'])
    doc = f"""<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><style>{CSS}{f.get('css','')}</style></head><body>
<header><div><span class="selo">REVISE PARA PROVA · SEU RESUMO EM UMA ÚNICA FOLHA</span><h1>{H.escape(f['titulo'])}</h1><p class="sub">{md(f['sub'])}</p></div><img src="file://{LOGO}" alt=""></header>
<div class="mais"><span>MAIS<br>COBRADO</span>{mais}</div>
<div class="grid"><div>{secs}</div><div><div class="lado"><div class="cab"><b>{H.escape(f['lado_tit'])}</b><span>{H.escape(f['lado_sub'])}</span></div><ol>{lado}</ol></div>
<div class="conta"><b class="t">{H.escape(f['conta_tit'])}</b>{f['conta_html']}</div></div></div>
<footer><span>{md(f['fontes'])}</span><b>MEDSUM · @medsum.oficial</b></footer></body></html>"""
    src = os.path.join(SAIDA, nome + '.html'); open(src, 'w').write(doc)
    pdf = os.path.join(SAIDA, nome + '.pdf'); png = os.path.join(SAIDA, nome + '.png')
    js = f"""const {{chromium}}=require('playwright');(async()=>{{const b=await chromium.launch({{executablePath:'/opt/pw-browsers/chromium'}});
const p=await b.newPage({{viewport:{{width:794,height:1123}},deviceScaleFactor:3}});await p.goto('file://{src}');await p.waitForTimeout(300);
const over=await p.evaluate(()=>{{const g=document.querySelector('.grid').getBoundingClientRect(), f=document.querySelector('footer').getBoundingClientRect(); return Math.round(g.bottom-f.top);}});
console.log('sobra(px, negativo=cabe):',over);
await p.pdf({{path:'{pdf}',format:'A4',printBackground:true,margin:{{top:0,right:0,bottom:0,left:0}}}});
await p.screenshot({{path:'{png}'}});await b.close();}})();"""
    jsf = os.path.join(SAIDA, nome + '.js'); open(jsf, 'w').write(js)
    env = dict(os.environ, NODE_PATH=subprocess.run(['npm', 'root', '-g'], capture_output=True, text=True).stdout.strip())
    out = subprocess.run(['node', jsf], capture_output=True, text=True, env=env); print(nome, out.stdout.strip(), out.stderr[-300:])
    os.remove(jsf)
    # Word: página inteira como imagem (igual à folha do Diabetes)
    import docx
    from docx.shared import Cm
    d = docx.Document(); s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    for a in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(s, a, Cm(0))
    s.header_distance = s.footer_distance = Cm(0)
    p = d.add_paragraph(); p.paragraph_format.space_after = 0; p.paragraph_format.space_before = 0
    p.add_run().add_picture(png, width=Cm(21.0))
    d.save(os.path.join(SAIDA, nome + '.docx'))
    n = subprocess.run(['pdfinfo', pdf], capture_output=True, text=True).stdout
    print('  páginas:', re.search(r'Pages:\s+(\d+)', n).group(1))
