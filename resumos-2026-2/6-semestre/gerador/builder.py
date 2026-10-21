# Gera um resumo MedSum em Word a partir do molde (Diabetes revisado 20/10/2026).
import copy, re, subprocess, os
import docx
from docx.shared import Cm
from docx.oxml.ns import qn
from lxml import etree

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
VINHO = '921827'


def el(tag, **attrs):
    e = etree.SubElement(etree.Element('dummy'), qn(tag))
    e = etree.Element(qn(tag))
    for k, v in attrs.items():
        e.set(qn(k), v)
    return e


def runs_from_markup(text, size=None, color=None, italic=False, base_bold=False):
    """**negrito**, __negrito sublinhado__, ^^sobrescrito^^ ."""
    out = []
    for part in re.split(r'(\*\*.+?\*\*|__.+?__|\^\^.+?\^\^)', text):
        if not part:
            continue
        bold = base_bold; und = False; sup = False
        if part.startswith('**'):
            part = part[2:-2]; bold = True
        elif part.startswith('__'):
            part = part[2:-2]; bold = True; und = True
        elif part.startswith('^^'):
            part = part[2:-2]; sup = True
        r = etree.Element(qn('w:r'))
        rpr = etree.SubElement(r, qn('w:rPr'))
        if bold:
            etree.SubElement(rpr, qn('w:b')); etree.SubElement(rpr, qn('w:bCs'))
        if italic:
            etree.SubElement(rpr, qn('w:i'))
        if color:
            c = etree.SubElement(rpr, qn('w:color')); c.set(qn('w:val'), color)
        if size:
            for t in ('w:sz', 'w:szCs'):
                s = etree.SubElement(rpr, qn(t)); s.set(qn('w:val'), str(size))
        if und:
            u = etree.SubElement(rpr, qn('w:u')); u.set(qn('w:val'), 'single')
        if sup:
            v = etree.SubElement(rpr, qn('w:vertAlign')); v.set(qn('w:val'), 'superscript')
        t = etree.SubElement(r, qn('w:t')); t.text = part
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        out.append(r)
    return out


def para(text='', style=None, ppr_xml=None, **kw):
    p = etree.Element(qn('w:p'))
    if ppr_xml is not None:
        p.append(copy.deepcopy(ppr_xml))
    elif style:
        ppr = etree.SubElement(p, qn('w:pPr'))
        ps = etree.SubElement(ppr, qn('w:pStyle')); ps.set(qn('w:val'), style)
    for r in runs_from_markup(text, **kw):
        p.append(r)
    return p


def ppr(spacing=None, jc=None, keep=False, num=None):
    x = etree.Element(qn('w:pPr'))
    if keep:
        etree.SubElement(x, qn('w:keepNext'))
    if num is not None:
        n = etree.SubElement(x, qn('w:numPr'))
        a = etree.SubElement(n, qn('w:ilvl')); a.set(qn('w:val'), str(num[1]))
        b = etree.SubElement(n, qn('w:numId')); b.set(qn('w:val'), str(num[0]))
    if spacing:
        s = etree.SubElement(x, qn('w:spacing'))
        for k, v in spacing.items():
            s.set(qn('w:' + k), str(v))
    if jc:
        j = etree.SubElement(x, qn('w:jc')); j.set(qn('w:val'), jc)
    return x


def cell(width, fill, text_lines, size=18, bold_first_col=False, header=False, left_bar=None, mar=(60, 90)):
    tc = etree.Element(qn('w:tc'))
    tcpr = etree.SubElement(tc, qn('w:tcPr'))
    w = etree.SubElement(tcpr, qn('w:tcW')); w.set(qn('w:w'), str(width)); w.set(qn('w:type'), 'dxa')
    bd = etree.SubElement(tcpr, qn('w:tcBorders'))
    for side in ('top', 'left', 'bottom', 'right'):
        s = etree.SubElement(bd, qn('w:' + side))
        if side == 'left' and left_bar:
            s.set(qn('w:val'), 'single'); s.set(qn('w:sz'), '24'); s.set(qn('w:space'), '0'); s.set(qn('w:color'), left_bar)
        elif side == 'bottom' and not header and not left_bar:
            s.set(qn('w:val'), 'single'); s.set(qn('w:sz'), '4'); s.set(qn('w:space'), '0'); s.set(qn('w:color'), 'E3C9CE')
        else:
            s.set(qn('w:val'), 'nil')
    sh = etree.SubElement(tcpr, qn('w:shd')); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), fill)
    m = etree.SubElement(tcpr, qn('w:tcMar'))
    for side, v in (('top', mar[0]), ('left', mar[1]), ('bottom', mar[0]), ('right', mar[1])):
        s = etree.SubElement(m, qn('w:' + side)); s.set(qn('w:w'), str(v)); s.set(qn('w:type'), 'dxa')
    va = etree.SubElement(tcpr, qn('w:vAlign')); va.set(qn('w:val'), 'center')
    if isinstance(text_lines, str):
        text_lines = [text_lines]
    for line in text_lines:
        bullet = line.startswith('• ')
        p = etree.SubElement(tc, qn('w:p'))
        p.append(ppr(spacing={'after': 0, 'line': 264, 'lineRule': 'auto'}, jc='left'))
        for r in runs_from_markup(line, size=size, color='FFFFFF' if header else None,
                                  base_bold=header or bold_first_col):
            p.append(r)
    return tc


def table(headers, rows, widths, size=18):
    tbl = etree.Element(qn('w:tbl'))
    tp = etree.SubElement(tbl, qn('w:tblPr'))
    tw = etree.SubElement(tp, qn('w:tblW')); tw.set(qn('w:w'), str(sum(widths))); tw.set(qn('w:type'), 'dxa')
    tl = etree.SubElement(tp, qn('w:tblLayout')); tl.set(qn('w:type'), 'fixed')
    g = etree.SubElement(tbl, qn('w:tblGrid'))
    for wd in widths:
        gc = etree.SubElement(g, qn('w:gridCol')); gc.set(qn('w:w'), str(wd))
    if headers:
        tr = etree.SubElement(tbl, qn('w:tr'))
        trp = etree.SubElement(tr, qn('w:trPr')); etree.SubElement(trp, qn('w:tblHeader'))
        for h, wd in zip(headers, widths):
            tr.append(cell(wd, VINHO, h, size=size, header=True))
    for i, row in enumerate(rows):
        tr = etree.SubElement(tbl, qn('w:tr'))
        trp = etree.SubElement(tr, qn('w:trPr')); etree.SubElement(trp, qn('w:cantSplit'))
        fill = 'FFFFFF' if i % 2 == 0 else 'F8EDEF'
        for j, (c, wd) in enumerate(zip(row, widths)):
            tr.append(cell(wd, fill, c, size=size, bold_first_col=(j == 0 and len(widths) > 1)))
    return tbl


class Resumo:
    def __init__(self, molde):
        self.doc = docx.Document(molde)
        self.body = self.doc.element.body
        b = list(self.body)
        self.proto_banner = copy.deepcopy(b[101])
        self.proto_h = {3: copy.deepcopy(b[102].find(qn('w:pPr'))), 4: copy.deepcopy(b[103].find(qn('w:pPr'))),
                        5: copy.deepcopy(b[113].find(qn('w:pPr')))}
        self.toc_first = b[9]
        self.toc_end = b[98]
        # apaga o sumário antigo (mantém o 1º e o fim do campo) e todo o corpo
        for e in b[10:98]:
            self.body.remove(e)
        for e in b[101:-1]:
            self.body.remove(e)
        self.sect = b[-1]
        self.n_tab = 0; self.n_fig = 0; self.n_quadro = 0
        self.headings = []  # (nivel, texto)

    def add(self, e):
        self.sect.addprevious(e)
        return e

    # ---------- capa ----------
    def capa(self, titulo, subtitulo, linha_fontes, capitulos, rodape):
        b = list(self.body)
        t = b[2]
        ts = [x for x in t.iter(qn('w:t'))]
        ts[0].text = titulo; ts[1].text = subtitulo; ts[2].text = linha_fontes
        for x in ts[3:]:
            x.text = ''
        grid = b[4]
        tcs = list(grid.iter(qn('w:tc')))
        for k, tc in enumerate(tcs):
            tt = list(tc.iter(qn('w:t')))
            if k < len(capitulos):
                tt[0].text = f'Capítulo {k + 1}'; tt[1].text = capitulos[k]
                for x in tt[2:]:
                    x.text = ''
            else:
                for x in tt:
                    x.text = ''
        # remove linhas vazias sobrando na grade de capítulos
        for tr in list(grid.iter(qn('w:tr'))):
            if not ''.join(x.text or '' for x in tr.iter(qn('w:t'))).strip():
                grid.remove(tr)
        tt = list(b[6].iter(qn('w:t')))
        tt[0].text = rodape
        for x in tt[1:]:
            x.text = ''

    # ---------- blocos ----------
    def faixa(self, titulo, sub):
        e = copy.deepcopy(self.proto_banner)
        ts = list(e.iter(qn('w:t')))
        for i, t in enumerate(ts):
            t.text = titulo if i % 2 == 0 else sub
        # a faixa começa em página nova (a 1ª já vem depois da quebra do sumário)
        if getattr(self, '_ja_teve_faixa', False):
            self.add(para('', ppr_xml=ppr(spacing={'after': 0, 'line': 240, 'lineRule': 'auto'}))).append(
                etree.fromstring(f'<w:r xmlns:w="{W}"><w:br w:type="page"/></w:r>'))
        self._ja_teve_faixa = True
        self.add(e)

    def h(self, nivel, texto):
        self.headings.append((nivel, texto))
        return self.add(para(texto, ppr_xml=self.proto_h[nivel]))

    def p(self, texto):
        return self.add(para(texto))

    def bullets(self, itens, num=1):
        for it in itens:
            self.add(para(it, ppr_xml=ppr(num=(num, 0))))

    def numerada(self, itens, num):
        for it in itens:
            self.add(para(it, ppr_xml=ppr(num=(num, 0))))

    def _legenda(self, rotulo, texto, size=18):
        p = para('', ppr_xml=ppr(spacing={'before': 60, 'after': 60, 'line': 240, 'lineRule': 'auto'}, jc='left', keep=True))
        for r in runs_from_markup(rotulo, size=size, color=VINHO, base_bold=True):
            p.append(r)
        for r in runs_from_markup(texto, size=size, color=VINHO):
            p.append(r)
        self.add(p)

    def fonte(self, texto):
        self.add(para(texto, ppr_xml=ppr(spacing={'before': 60, 'line': 240, 'lineRule': 'auto'}),
                      size=16, color='595959', italic=True))

    def tabela(self, titulo, headers, rows, widths, fonte, quadro=False):
        if quadro:
            self.n_quadro += 1; rot = f'Quadro {self.n_quadro}'
        else:
            self.n_tab += 1; rot = f'Tabela {self.n_tab}'
        self._legenda(rot, ' – ' + titulo)
        self.add(table(headers, rows, widths))
        if fonte:
            self.fonte(fonte)

    def caixa(self, titulo, linhas, fill='F8EDEF', barra=VINHO):
        tbl = etree.Element(qn('w:tbl'))
        tp = etree.SubElement(tbl, qn('w:tblPr'))
        tw = etree.SubElement(tp, qn('w:tblW')); tw.set(qn('w:w'), '9900'); tw.set(qn('w:type'), 'dxa')
        tl = etree.SubElement(tp, qn('w:tblLayout')); tl.set(qn('w:type'), 'fixed')
        g = etree.SubElement(tbl, qn('w:tblGrid')); gc = etree.SubElement(g, qn('w:gridCol')); gc.set(qn('w:w'), '9900')
        tr = etree.SubElement(tbl, qn('w:tr'))
        trp = etree.SubElement(tr, qn('w:trPr')); etree.SubElement(trp, qn('w:cantSplit'))
        tc = cell(9900, fill, [], left_bar=barra, mar=(100, 160))
        p = etree.SubElement(tc, qn('w:p'))
        p.append(ppr(spacing={'after': 20, 'line': 276, 'lineRule': 'auto'}, jc='left', keep=True))
        for r in runs_from_markup(titulo, size=16, color=barra, base_bold=True):
            p.append(r)
        if isinstance(linhas, str):
            linhas = [linhas]
        for ln in linhas:
            p = etree.SubElement(tc, qn('w:p'))
            if ln.startswith('• '):
                p.append(ppr(num=(1, 0), spacing={'after': 20, 'line': 276, 'lineRule': 'auto'}, jc='left'))
                ln = ln[2:]
            else:
                p.append(ppr(spacing={'after': 40, 'line': 276, 'lineRule': 'auto'}))
            for r in runs_from_markup(ln, size=16):
                p.append(r)
        tr.append(tc)
        self.add(tbl)
        self.add(para(''))

    def ponto_chave(self, linhas):
        self.caixa('Ponto-chave', linhas)

    def caso(self, linhas):
        self.caixa('Aplicação aos casos da SP', linhas, fill='FBF4E6', barra='A8771C')

    def figura(self, png, titulo, fonte, largura_cm=16.5):
        self.n_fig += 1
        p = self.doc.add_paragraph()
        p.alignment = 1
        p.add_run().add_picture(png, width=Cm(largura_cm))
        self.body.remove(p._p)
        self.add(p._p)
        p._p.insert(0, ppr(jc='center', keep=False))
        self._legenda(f'Figura {self.n_fig}', ': ' + titulo, size=16)
        if fonte:
            self.fonte(fonte)

    def referencias(self, refs):
        self.h(3, 'Referências')
        for r in refs:
            pp = ppr(num=(84, 0), spacing={'line': 276, 'lineRule': 'auto'})
            self.add(para(r, ppr_xml=pp, size=20))

    # ---------- sumário ----------
    def sumario(self, paginas):
        """paginas: lista de números de página, na ordem self.headings (só níveis 3 e 4)."""
        ents = [(n, t) for n, t in self.headings if n in (3, 4)]
        first = self.toc_first
        # o 1º parágrafo guarda o início do campo TOC: troca só o texto e a página
        hl = first.find(qn('w:hyperlink'))
        ts = list(hl.iter(qn('w:t')))
        ts[0].text = ents[0][1]; ts[1].text = str(paginas[0])
        anchor = first
        for (n, t), pg in zip(ents[1:], paginas[1:]):
            p = etree.Element(qn('w:p'))
            pp = etree.SubElement(p, qn('w:pPr'))
            ps = etree.SubElement(pp, qn('w:pStyle')); ps.set(qn('w:val'), 'Sumrio3' if n == 3 else 'Sumrio4')
            tabs = etree.SubElement(pp, qn('w:tabs')); tb = etree.SubElement(tabs, qn('w:tab'))
            tb.set(qn('w:val'), 'right'); tb.set(qn('w:leader'), 'dot'); tb.set(qn('w:pos'), '9968')
            for r in runs_from_markup(t):
                p.append(r)
            p.append(etree.fromstring(f'<w:r xmlns:w="{W}"><w:tab/></w:r>'))
            p.append(etree.fromstring(f'<w:r xmlns:w="{W}"><w:t>{pg}</w:t></w:r>'))
            anchor.addnext(p); anchor = p

    def limpar_midias(self):
        part = self.doc.part
        xml = etree.tostring(self.body).decode()
        for hdr in [s for s in self.doc.sections]:
            pass
        for rid, rel in list(part.rels.items()):
            if 'image' in rel.reltype and f'"{rid}"' not in xml:
                del part.rels[rid]

    def salvar(self, caminho):
        self.limpar_midias()
        self.doc.save(caminho)


def paginas_dos_titulos(pdf, headings, pagina_inicial_corpo=3):
    """Acha a página de cada título (níveis 3 e 4) no PDF, em ordem."""
    txt = subprocess.run(['pdftotext', '-layout', pdf, '-'], capture_output=True, text=True).stdout
    pages = txt.split('\f')
    norm = lambda s: re.sub(r'\s+', ' ', s).strip().lower()
    pnorm = [norm(p) for p in pages]
    res = []; cur = pagina_inicial_corpo - 1
    for n, t in headings:
        if n not in (3, 4):
            continue
        alvo = norm(t)[:40]
        k = cur
        while k < len(pnorm) and alvo not in pnorm[k]:
            k += 1
        if k >= len(pnorm):
            k = cur
        cur = k
        res.append(k + 1)
    return res
